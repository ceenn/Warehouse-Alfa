"""Laporan: builder data + exporter Excel (openpyxl) dan PDF (reportlab)."""
from datetime import timezone, timedelta, datetime
from io import BytesIO
from xml.sax.saxutils import escape
from .core import AppError
from .database import Database

WIB = timezone(timedelta(hours=7))


def fdate(dt):
    return dt.astimezone(WIB).strftime("%d-%m-%Y %H:%M") if dt else ""


class ReportBuilder:
    TITLES = {
        "inventory": "Stok & Saldo Inventory",
        "low_stock": "Barang di Bawah Stok Minimum",
        "stock_card": "Kartu Stok (Transaksi IN / OUT / ADJ)",
        "requests": "Rekap Request Barang",
        "purchase_orders": "Rekap Purchase Order",
    }

    @classmethod
    def build(cls, kind):
        if kind not in cls.TITLES:
            raise AppError("Jenis laporan tidak dikenal", 404)
        headers, rows = getattr(cls, "_" + kind)()
        return {"kind": kind, "title": cls.TITLES[kind], "headers": headers, "rows": rows}

    @staticmethod
    def _inventory():
        h = ["Kode", "Nama Barang", "Kategori", "Warehouse", "Lokasi", "Stok", "Min", "Satuan", "Status"]
        rows = [[i["code"], i["name"], i["category"], i.get("warehouse", ""), i.get("location", ""), i.get("stock", 0),
                 i.get("min_stock", 0), i.get("unit", ""), "Menipis" if i.get("stock", 0) <= i.get("min_stock", 0) else "Aman"]
                for i in Database.col("items").find().sort("code", 1)]
        return h, rows

    @classmethod
    def _low_stock(cls):
        h, rows = cls._inventory()
        return h, [r for r in rows if r[-1] == "Menipis"]

    @staticmethod
    def _stock_card():
        h = ["Tanggal", "Kode", "Barang", "Tipe", "Qty", "Saldo", "Referensi", "Oleh"]
        rows = [[fdate(t["at"]), t["code"], t["name"], t["type"], t["qty"], t["balance"], t.get("ref_no", ""), t.get("by", "")]
                for t in Database.col("stock_transactions").find().sort("at", -1).limit(2000)]
        return h, rows

    @staticmethod
    def _requests():
        h = ["No", "Pemohon", "Barang", "Status", "Tanggal", "Approver", "Diserahkan oleh"]
        rows = [[r["number"], r["requester"], ", ".join(f"{l['name']} x{l['qty']}" for l in r["lines"]), r["status"],
                 fdate(r["created_at"]), r.get("approver", ""), r.get("fulfilled_by", "")]
                for r in Database.col("requests").find().sort("created_at", -1)]
        return h, rows

    @staticmethod
    def _purchase_orders():
        h = ["No", "Supplier", "Barang (diterima/dipesan)", "Total (Rp)", "Status", "Tanggal", "Dibuat oleh", "Approver"]
        rows = [[p["number"], p["supplier"], ", ".join(f"{l['name']} {l.get('received', 0)}/{l['qty']}" for l in p["lines"]),
                 p.get("total", 0), p["status"], fdate(p["created_at"]), p["creator"], p.get("approver", "")]
                for p in Database.col("purchase_orders").find().sort("created_at", -1)]
        return h, rows


class ExcelExporter:
    mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ext = "xlsx"

    @staticmethod
    def export(rep):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = rep["title"][:30].replace("/", "-")
        ws.append([rep["title"]])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"Dicetak: {datetime.now(WIB).strftime('%d-%m-%Y %H:%M')} WIB"])
        ws.append([])
        ws.append(rep["headers"])
        for c in ws[4]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F6FEB")
        for r in rep["rows"]:
            ws.append(r)
        for idx, h in enumerate(rep["headers"], 1):
            mx = max([len(str(h))] + [len(str(r[idx - 1])) for r in rep["rows"]])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(mx + 2, 10), 60)
        ws.freeze_panes = "A5"
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf


class PdfExporter:
    mimetype = "application/pdf"
    ext = "pdf"

    @staticmethod
    def export(rep):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        buf = BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
        ss = getSampleStyleSheet()
        cell = ParagraphStyle("cell", parent=ss["BodyText"], fontSize=8, leading=10)
        head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
        data = [[Paragraph(escape(str(h)), head) for h in rep["headers"]]]
        data += [[Paragraph(escape(str(c)), cell) for c in r] for r in rep["rows"]]
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F6FEB")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0B7C3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FB")]),
        ]))
        doc.build([Paragraph(escape(rep["title"]), ss["Title"]),
                   Paragraph(f"Dicetak: {datetime.now(WIB).strftime('%d-%m-%Y %H:%M')} WIB", ss["Normal"]),
                   Spacer(1, 10), table])
        buf.seek(0)
        return buf


EXPORTERS = {"xlsx": ExcelExporter, "pdf": PdfExporter}
